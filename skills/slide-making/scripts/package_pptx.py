#!/usr/bin/env python3
"""把 Beamer PDF 渲出的整页 PNG 逐页贴回官方 pptx 模板（slide-making §6.1 交付包装）。

固化 SKILL §6.1 全流程，含 §7 协作红线：
  1. 锁文件检查——目录里有 `~$xxx.pptx` 说明 PowerPoint 正开着该文件，拒绝写入
  2. 自动备份官方模板（同目录，带时间戳，原文件绝不改动）
  3. 删除模板全部示例页
  4. 选占位符最少的版式（空白），清掉新页上的残留占位符
  5. 逐页贴整幅图片于 (0,0) 满幅
  6. 保存输出文件 → 重新打开验证页数与图片位置

用法：
    python package_pptx.py <官方模板.pptx> <PNG目录> -o <输出.pptx>

    pdftoppm -png -r 300 presentation.pdf /tmp/slide-png
    python package_pptx.py official.pptx /tmp/slide-png -o packaged.pptx

原图与模板必须同为 16:9，否则贴图变形，脚本会校验并中止。
依赖 python-pptx 与 Pillow：
    uv run --with python-pptx --with pillow python package_pptx.py ...
"""
import argparse
import re
import shutil
import sys
import time
from pathlib import Path

from PIL import Image
from pptx import Presentation

RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def natural_key(path: Path):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", path.name)]


def check_lock_file(target: Path):
    locks = sorted(target.parent.glob("~$*"))
    if locks:
        sys.exit("package_pptx: 发现 PowerPoint 锁文件 "
                 f"{', '.join(l.name for l in locks)}——目标文件正被 PowerPoint 打开，"
                 "写入会被旧窗口重存冲掉。请先关闭 PowerPoint 再运行。")


def check_pngs(pngs: list[Path]):
    """整页渲染图必然全部同尺寸；混入其他图（如 contact sheet）说明目录不纯，拒绝包装。"""
    dims = [Image.open(p).size for p in pngs]
    if len(set(dims)) > 1:
        detail = ", ".join(f"{p.name}:{w}x{h}" for p, (w, h) in zip(pngs, dims))
        sys.exit("package_pptx: PNG 尺寸不齐，目录疑似混入非整页渲染图：\n  " + detail)


def check_aspect_ratio(prs: Presentation, pngs: list[Path]):
    img_w, img_h = Image.open(pngs[0]).size
    slide_ar = prs.slide_width / prs.slide_height
    img_ar = img_w / img_h
    if abs(slide_ar - img_ar) > 0.01:
        sys.exit(f"package_pptx: 宽高比不一致——模板 {slide_ar:.4f}（16:9=1.7778），"
                 f"图片 {img_ar:.4f}（{img_w}x{img_h}）。强行贴图会变形，"
                 "请检查 pdftoppm 出图与 pptx 页面设置是否同为 16:9。")


def delete_all_slides(prs: Presentation):
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        prs.part.drop_rel(sld.get(RID))
        xml_slides.remove(sld)


def add_full_bleed_slide(prs: Presentation, png: Path):
    layout = min(prs.slide_layouts, key=lambda l: len(l.placeholders))
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    slide.shapes.add_picture(str(png), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)


def verify_output(path: Path, n_slides: int):
    prs = Presentation(path)
    problems = []
    actual = len(prs.slides)
    if actual != n_slides:
        problems.append(f"页数不符：期望 {n_slides}，实际 {actual}")
    for i, slide in enumerate(prs.slides, 1):
        pics = [s for s in slide.shapes if s.shape_type == 13]  # PICTURE
        if len(pics) != 1:
            problems.append(f"第 {i} 页图片数 {len(pics)} != 1")
            continue
        pic = pics[0]
        if (pic.left, pic.top, pic.width, pic.height) != (
                0, 0, prs.slide_width, prs.slide_height):
            problems.append(f"第 {i} 页图片未满幅于 (0,0)")
    if problems:
        sys.exit("package_pptx: 复核失败：\n  " + "\n  ".join(problems))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("template", help="会议官方模板 pptx")
    ap.add_argument("png_dir", help="pdftoppm 出图目录（整页 PNG，按页码排序）")
    ap.add_argument("-o", "--output", help="输出 pptx，默认 <模板名>-packaged.pptx")
    args = ap.parse_args()

    template = Path(args.template)
    if not template.is_file():
        sys.exit(f"package_pptx: 模板不存在：{template}")
    pngs = sorted(Path(args.png_dir).glob("*.png"), key=natural_key)
    if not pngs:
        sys.exit(f"package_pptx: 没有找到 PNG：{args.png_dir}")
    output = Path(args.output) if args.output else \
        template.with_name(f"{template.stem}-packaged.pptx")

    check_lock_file(template)
    check_lock_file(output)

    backup = template.with_name(
        f"{template.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{template.suffix}")
    shutil.copy2(template, backup)

    prs = Presentation(str(template))
    check_pngs(pngs)
    check_aspect_ratio(prs, pngs)
    delete_all_slides(prs)
    for png in pngs:
        add_full_bleed_slide(prs, png)
    prs.save(str(output))

    verify_output(output, len(pngs))
    print(f"package_pptx: {len(pngs)} 页已贴入 → {output}\n"
          f"package_pptx: 原模板备份于 {backup}（未改动）")


if __name__ == "__main__":
    main()
